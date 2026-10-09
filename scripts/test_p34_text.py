"""p34 para_text_full: what Word stores as elements must appear in the text."""
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

sys.path.insert(0, str(Path(__file__).resolve().parent))
import p60_monthly_update as p60  # noqa: E402

p34 = p60._load_module(p60.P34_PATH, "p34_rebuild_from_hudoc")


def paragraph(*parts):
    """A paragraph built from text runs and bare elements ("noBreakHyphen", "softHyphen", "tab")."""
    p = Document().add_paragraph()
    for part in parts:
        if part in ("noBreakHyphen", "softHyphen", "tab"):
            run = p.add_run()
            run._r.append(OxmlElement(f"w:{part}"))
        else:
            p.add_run(part)
    return p


class TextTests(unittest.TestCase):
    def test_non_breaking_hyphen_is_kept(self):
        self.assertEqual(p34.para_text_full(paragraph("in respect of non", "noBreakHyphen", "pecuniary damage")),
                         "in respect of non-pecuniary damage")

    def test_soft_hyphen_stays_invisible(self):
        self.assertEqual(p34.para_text_full(paragraph("demo", "softHyphen", "cratic")), "democratic")

    def test_tab(self):
        self.assertEqual(p34.para_text_full(paragraph("12.", "tab", "Text")), "12.\tText")


if __name__ == "__main__":
    unittest.main()
